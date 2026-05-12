"""
python fix_excel_upload.py
سیستم آپلود محصولات با Excel از پنل ادمین
"""
import os, subprocess, sys

print("ساخت سیستم آپلود Excel ...")

# ─── 1. نصب openpyxl ─────────────────────────────────────
r = subprocess.run([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'],
                   capture_output=True, text=True)
print("  OK - openpyxl نصب شد")

# ─── 2. management command ───────────────────────────────
cmd_dir = os.path.join('apps','catalog','management','commands')
os.makedirs(cmd_dir, exist_ok=True)

cmd_content = '''"""
python manage.py import_products --file products.xlsx
"""
import os
from django.core.management.base import BaseCommand
from apps.catalog.excel_importer import import_products_from_excel


class Command(BaseCommand):
    help = "وارد کردن محصولات از فایل Excel"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help="مسیر فایل Excel")
        parser.add_argument("--dry-run", action="store_true", help="فقط نمایش بدون ذخیره")

    def handle(self, *args, **options):
        path = options["file"]
        if not os.path.exists(path):
            self.stderr.write(f"فایل یافت نشد: {path}")
            return

        results = import_products_from_excel(path, dry_run=options["dry_run"])
        self.stdout.write(f"\\n✓ اضافه: {results['created']}")
        self.stdout.write(f"✓ بروزرسانی: {results['updated']}")
        if results["errors"]:
            self.stdout.write(f"✗ خطا: {len(results['errors'])}")
            for e in results["errors"][:5]:
                self.stdout.write(f"  - {e}")
'''
with open(os.path.join(cmd_dir, 'import_products.py'), 'w', encoding='utf-8') as f:
    f.write(cmd_content)
print("  OK - management command")

# ─── 3. excel_importer.py ────────────────────────────────
importer_content = '''"""
apps/catalog/excel_importer.py
وارد کردن محصولات از Excel
"""
import openpyxl
from django.db import transaction
from apps.catalog.models import Category, Product, ProductVariant


# نگاشت نام ستون‌ها
COLUMNS = {
    # اطلاعات اصلی
    "نام محصول":       "name",
    "دسته بندی":       "category_slug",
    "شرکت سازنده":     "brand",
    "شماره مدل":       "model_number",
    "توضیح کوتاه":     "short_description",
    "توضیحات":         "description",
    # قیمت
    "قیمت خرید (یوان)":  "purchase_price_cny",
    "هزینه حمل (یوان)":  "shipping_cost_cny",
    "سود (%)":            "profit_percent",
    "موجودی":             "stock_quantity",
    "کد SKU":             "sku",
    # برد
    "نام برد":          "board_name",
    "سیستم":            "board_system",
    "رزولوشن برد":      "board_resolution",
    "نام سنسور":        "sensor_name",
    "اندازه سنسور":     "sensor_size",
    "AI پردازنده":      "ai_processor",
    "ابعاد برد":        "board_dimensions",
    # لنز
    "نام لنز":          "lens_name",
    "فاصله کانونی":     "lens_degree",
    "نوع لنز":          "lens_type",
    "نوع کاربرد":       "focal_ratio",
    "رزولوشن لنز":      "lens_resolution",
    # متعلقات
    "مدل":              "modem_model",
    "نوع":              "modem_type",
    "تعداد پین":        "pin_count",
    "ولتاژ":            "voltage",
    "آمپر خروجی":       "ampere_output",
    "قابلیت سیم کارت": "sim_card",
    "تعداد پورت":       "output_ports",
    "سایز":             "modem_size",
    # بدنه
    "جنس بدنه":         "housing_material",
    "نوع بدنه":         "housing_type",
    "سایز دوربین":      "housing_size",
    "قابلیت میکروفون": "has_microphone",
}

BOOL_FIELDS = {"sim_card", "has_microphone"}
INT_FIELDS  = {"stock_quantity", "output_ports"}
DEC_FIELDS  = {"purchase_price_cny", "shipping_cost_cny", "profit_percent"}


def import_products_from_excel(file_path: str, dry_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # پیدا کردن sheet محصولات
    ws = None
    for name in ["محصولات", "Products", "products", wb.sheetnames[0]]:
        if name in wb.sheetnames:
            ws = wb[name]
            break
    if not ws:
        ws = wb.active

    results = {"created": 0, "updated": 0, "errors": []}

    # خواندن هدر
    headers = {}
    for col in ws.iter_cols(1, ws.max_column, 1, 1):
        cell = col[0]
        if cell.value and str(cell.value).strip() in COLUMNS:
            headers[cell.column] = COLUMNS[str(cell.value).strip()]

    if not headers:
        results["errors"].append("ستون‌های مورد نظر در فایل پیدا نشد")
        return results

    # cache دسته‌بندی‌ها
    cat_cache = {c.slug: c for c in Category.objects.filter(is_active=True)}
    cat_name_cache = {c.name: c for c in Category.objects.filter(is_active=True)}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        # ساخت dict از ردیف
        data = {}
        for col_idx, field in headers.items():
            val = row[col_idx - 1]
            if val is not None:
                data[field] = val

        if not data.get("name"):
            continue

        try:
            with transaction.atomic():
                # پیدا کردن دسته‌بندی
                cat = None
                cat_val = str(data.pop("category_slug", "") or "").strip()
                if cat_val:
                    cat = cat_cache.get(cat_val) or cat_name_cache.get(cat_val)

                if not cat:
                    results["errors"].append(f"ردیف {row_idx}: دسته‌بندی '{cat_val}' پیدا نشد")
                    continue

                # اطلاعات variant
                sku               = str(data.pop("sku", "") or "").strip()
                purchase_price_cny= data.pop("purchase_price_cny", 0) or 0
                shipping_cost_cny = data.pop("shipping_cost_cny", 0) or 0
                profit_percent    = data.pop("profit_percent", 0) or 0
                stock_quantity    = int(data.pop("stock_quantity", 0) or 0)

                # تبدیل نوع‌ها
                for field in list(data.keys()):
                    val = data[field]
                    if field in BOOL_FIELDS:
                        data[field] = str(val).strip() in ("1","بله","yes","true","دارد","✓")
                    elif field in INT_FIELDS:
                        try: data[field] = int(val)
                        except: data[field] = None
                    elif field in DEC_FIELDS:
                        try: data[field] = float(val)
                        except: data[field] = 0
                    elif val is not None:
                        data[field] = str(val).strip()

                # ساخت یا بروزرسانی محصول
                if not dry_run:
                    product, created = Product.objects.update_or_create(
                        name=data["name"],
                        defaults={**data, "category": cat},
                    )
                    if created:
                        results["created"] += 1
                    else:
                        results["updated"] += 1

                    # ساخت یا بروزرسانی variant
                    if not sku:
                        sku = f"AUTO-{product.pk}"
                    ProductVariant.objects.update_or_create(
                        sku=sku,
                        defaults={
                            "product": product,
                            "purchase_price_cny": purchase_price_cny,
                            "shipping_cost_cny":  shipping_cost_cny,
                            "profit_percent":     profit_percent,
                            "stock_quantity":     stock_quantity,
                            "is_active": True,
                        }
                    )
                else:
                    results["created"] += 1

        except Exception as e:
            results["errors"].append(f"ردیف {row_idx}: {e}")

    return results
'''
with open(os.path.join('apps','catalog','excel_importer.py'), 'w', encoding='utf-8') as f:
    f.write(importer_content)
print("  OK - excel_importer.py")

# ─── 4. admin: اضافه کردن آپلود Excel ───────────────────
admin_path = os.path.join('apps','catalog','admin.py')
with open(admin_path, encoding='utf-8') as f:
    content = f.read()

# اضافه کردن import ها
old_imports = 'from django.contrib import admin\nfrom django.utils.html import format_html'
new_imports = '''from django.contrib import admin
from django.utils.html import format_html
from django.urls import path
from django.shortcuts import render, redirect
from django.contrib import messages
from django import forms


class ExcelUploadForm(forms.Form):
    excel_file = forms.FileField(label="فایل Excel", help_text="فرمت xlsx")'''

content = content.replace(old_imports, new_imports)

# اضافه کردن دکمه و view به ProductAdmin
old_admin = '''@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    list_display    = ["name", "category", "is_active", "is_featured"]'''

new_admin = '''@admin.register(Product)
class ProductAdmin(admin.ModelAdmin):
    list_display    = ["name", "category", "is_active", "is_featured"]
    change_list_template = "admin/catalog/product/change_list.html"

    def get_urls(self):
        urls = super().get_urls()
        extra = [
            path("import-excel/",
                 self.admin_site.admin_view(self.import_excel_view),
                 name="catalog_product_import_excel"),
        ]
        return extra + urls

    def import_excel_view(self, request):
        from apps.catalog.excel_importer import import_products_from_excel
        import tempfile, os
        if request.method == "POST":
            form = ExcelUploadForm(request.POST, request.FILES)
            if form.is_valid():
                f = request.FILES["excel_file"]
                # ذخیره موقت
                with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
                    for chunk in f.chunks():
                        tmp.write(chunk)
                    tmp_path = tmp.name
                try:
                    dry_run = "dry_run" in request.POST
                    results = import_products_from_excel(tmp_path, dry_run=dry_run)
                    if dry_run:
                        messages.info(request, f"[پیش‌نمایش] {results['created']} محصول اضافه می‌شد | {results['updated']} بروزرسانی")
                    else:
                        messages.success(request, f"✓ {results['created']} محصول اضافه شد | {results['updated']} بروزرسانی")
                    if results["errors"]:
                        for e in results["errors"][:5]:
                            messages.warning(request, e)
                except Exception as e:
                    messages.error(request, f"خطا: {e}")
                finally:
                    os.unlink(tmp_path)
                return redirect("..")
        else:
            form = ExcelUploadForm()
        return render(request, "admin/catalog/product/import_excel.html", {
            "form": form,
            "title": "آپلود محصولات از Excel",
            "opts": self.model._meta,
        })'''

content = content.replace(old_admin, new_admin)
with open(admin_path, 'w', encoding='utf-8') as f:
    f.write(content)
print("  OK - admin.py")

# ─── 5. template های admin ───────────────────────────────
tpl_dir = os.path.join('templates','admin','catalog','product')
os.makedirs(tpl_dir, exist_ok=True)

# change_list.html
with open(os.path.join(tpl_dir,'change_list.html'), 'w', encoding='utf-8') as f:
    f.write('''{% extends "admin/change_list.html" %}
{% block object-tools-items %}
<li>
  <a href="import-excel/" class="addlink" style="background:#417690">
    📊 آپلود Excel
  </a>
</li>
{{ block.super }}
{% endblock %}
''')

# import_excel.html
with open(os.path.join(tpl_dir,'import_excel.html'), 'w', encoding='utf-8') as f:
    f.write('''{% extends "admin/base_site.html" %}
{% block content %}
<div style="direction:rtl;max-width:600px;margin:20px auto">
  <h2>📊 آپلود محصولات از Excel</h2>
  <div style="background:#e8f4f8;padding:15px;border-radius:8px;margin-bottom:20px">
    <b>راهنما:</b>
    <ul style="margin:10px 0 0 20px">
      <li>فایل باید فرمت <b>.xlsx</b> باشد</li>
      <li>Sheet اول یا sheet با نام «محصولات» خوانده می‌شود</li>
      <li>ردیف اول باید نام ستون‌ها باشد</li>
      <li>ستون «نام محصول» و «دسته بندی» اجباری هستند</li>
      <li>اگر محصول با همین نام وجود داشته باشد بروزرسانی می‌شود</li>
    </ul>
  </div>
  <form method="post" enctype="multipart/form-data" style="background:#fff;padding:20px;border-radius:8px;border:1px solid #ddd">
    {% csrf_token %}
    <table>{{ form.as_table }}</table>
    <br>
    <label style="display:flex;align-items:center;gap:8px;margin-bottom:15px">
      <input type="checkbox" name="dry_run"> پیش‌نمایش (بدون ذخیره)
    </label>
    <button type="submit" class="default" style="padding:8px 20px">آپلود و وارد کردن</button>
    <a href=".." style="margin-right:10px">انصراف</a>
  </form>
</div>
{% endblock %}
''')
print("  OK - templates")

# ─── 6. settings: اضافه کردن templates dir ──────────────
settings_path = os.path.join('config','settings.py')
with open(settings_path, encoding='utf-8') as f:
    settings = f.read()

if "'DIRS': []" in settings:
    settings = settings.replace(
        "'DIRS': []",
        "'DIRS': [BASE_DIR / 'templates']"
    )
    with open(settings_path, 'w', encoding='utf-8') as f:
        f.write(settings)
    print("  OK - settings.py")

# ─── 7. check ────────────────────────────────────────────
import ast
files = [
    os.path.join('apps','catalog','admin.py'),
    os.path.join('apps','catalog','excel_importer.py'),
]
all_ok = True
for fp in files:
    try:
        ast.parse(open(fp, encoding='utf-8').read())
        print(f"  OK syntax: {os.path.basename(fp)}")
    except SyntaxError as e:
        print(f"  ERROR: {fp}: {e}")
        all_ok = False

print("\n" + "="*55)
if all_ok:
    print("  تمام! حالا بزنید: python manage.py runserver")
    print("  بعد در پنل ادمین → محصولات → دکمه 'آپلود Excel'")
print("="*55)
