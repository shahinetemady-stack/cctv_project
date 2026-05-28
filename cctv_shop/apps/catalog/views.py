"""
apps/catalog/views.py
──────────────────────
Public API endpoints for the catalog.

GET  /api/catalog/categories/          لیست دسته‌بندی‌ها
GET  /api/catalog/products/            لیست محصولات (با فیلتر)
GET  /api/catalog/products/<slug>/     جزئیات محصول
GET  /api/catalog/products/search/     جستجوی هوشمند
GET  /api/catalog/exchange-rate/       نرخ ارز جاری
PUT  /api/catalog/exchange-rate/       بروزرسانی نرخ (فقط ادمین)
POST /api/catalog/products/bulk-upload/ آپلود اکسل (فقط ادمین)
"""

import logging
from django.db.models import Prefetch
from rest_framework import generics, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAdminUser, IsAuthenticatedOrReadOnly
from rest_framework.response import Response
from rest_framework.views import APIView

from .filters import ProductFilter
from .models import Category, ExchangeRateSetting, Product, ProductImage, ProductVariant
from .serializers import (
    CategorySerializer,
    ExchangeRateSerializer,
    ExchangeRateUpdateSerializer,
    ProductDetailSerializer,
    ProductListSerializer,
)

logger = logging.getLogger(__name__)


# ─── Categories ───────────────────────────────────────────────────────────────

class CategoryListView(generics.ListAPIView):
    """
    GET /api/catalog/categories/
    فقط دسته‌های ریشه (parent=None) با فرزندانشان.
    """
    serializer_class = CategorySerializer
    permission_classes = [AllowAny]

    def get_queryset(self):
        return (
            Category.objects.filter(is_active=True, parent__isnull=True)
            .prefetch_related("children")
            .order_by("sort_order", "name")
        )


# ─── Products ─────────────────────────────────────────────────────────────────

class ProductListView(generics.ListAPIView):
    """
    GET /api/catalog/products/
    پارامترهای فیلتر: part_type, chipset, voltage, lens_mount,
                       category, in_stock, is_featured, search
    مرتب‌سازی: ?ordering=name / -name / created_at / -created_at
    """
    serializer_class = ProductListSerializer
    permission_classes = [AllowAny]
    filterset_class = ProductFilter
    search_fields = ["name", "brand", "model_number", "chipset"]
    ordering_fields = ["name", "created_at", "is_featured"]
    ordering = ["-is_featured", "-created_at"]

    def list(self, request, *args, **kwargs):
        import math
        response = super().list(request, *args, **kwargs)
        if request.user.is_authenticated:
            try:
                discount = float(request.user.profile.discount_percent or 0)
                if discount > 0:
                    for product in response.data.get("results", response.data if isinstance(response.data, list) else []):
                        mp = product.get("min_price")
                        if mp:
                            product["partner_min_price"] = math.ceil(float(mp) * (1 - discount/100))
                            product["partner_discount_percent"] = discount
            except Exception:
                pass
        return response

    def get_queryset(self):
        qs = (
            Product.objects.filter(is_active=True)
            .select_related("category")
            .prefetch_related(
                Prefetch("images", queryset=ProductImage.objects.filter(is_primary=True)),
                Prefetch("variants", queryset=ProductVariant.objects.filter(is_active=True)),
            )
        )
        p = self.request.query_params

        # فیلترهای اختصاصی برد
        if p.get("sensor_name"):   qs = qs.filter(sensor_name__icontains=p["sensor_name"])
        if p.get("board_system"):  qs = qs.filter(board_system__icontains=p["board_system"])
        if p.get("board_resolution"): qs = qs.filter(board_resolution__icontains=p["board_resolution"])
        if p.get("sensor_size"):   qs = qs.filter(sensor_size__icontains=p["sensor_size"])

        # فیلترهای اختصاصی لنز
        if p.get("lens_degree"):   qs = qs.filter(lens_degree__icontains=p["lens_degree"])
        if p.get("lens_type"):     qs = qs.filter(lens_type__icontains=p["lens_type"])
        if p.get("focal_ratio"):   qs = qs.filter(focal_ratio__icontains=p["focal_ratio"])

        # فیلترهای اختصاصی متعلقات
        if p.get("modem_type"):    qs = qs.filter(modem_type__icontains=p["modem_type"])
        if p.get("voltage"):       qs = qs.filter(voltage__icontains=p["voltage"])
        if p.get("sim_card") == "true": qs = qs.filter(sim_card=True)

        # فیلترهای اختصاصی بدنه
        if p.get("housing_material"): qs = qs.filter(housing_material__icontains=p["housing_material"])
        if p.get("housing_type"):  qs = qs.filter(housing_type__icontains=p["housing_type"])
        if p.get("has_microphone") == "true": qs = qs.filter(has_microphone=True)

        return qs


class ProductDetailView(generics.RetrieveAPIView):
    """
    GET /api/catalog/products/<slug>/
    """
    serializer_class = ProductDetailSerializer
    permission_classes = [AllowAny]
    lookup_field = "slug"

    def get_queryset(self):
        return (
            Product.objects.filter(is_active=True)
            .select_related("category")
            .prefetch_related(
                "images",
                "variants",
                "compatible_with",
            )
        )

    def retrieve(self, request, *args, **kwargs):
        import math
        response = super().retrieve(request, *args, **kwargs)
        if request.user.is_authenticated:
            try:
                discount = float(request.user.profile.discount_percent or 0)
                if discount > 0:
                    for v in response.data.get("variants", []):
                        base = v.get("final_price", 0)
                        v["partner_price"] = math.ceil(base * (1 - discount/100))
                        v["partner_discount_percent"] = discount
            except Exception:
                pass
        return response


# ─── Smart Search ──────────────────────────────────────────────────────────────

class ProductSearchView(APIView):
    """
    GET /api/catalog/products/search/?q=<query>

    جستجوی هوشمند با اصلاح غلط‌های رایج فارسی:
    - اشتباهات رایج حروف (ک/ک، ی/ی، و/و ...)
    - جستجو در نام، برند، مدل، چیپست
    """
    permission_classes = [AllowAny]

    # نگاشت غلط‌های رایج فارسی (ساده‌ترین رویکرد بدون نیاز به کتابخانه خارجی)
    PERSIAN_CORRECTIONS = {
        "ك": "ک",   # عربی → فارسی
        "ي": "ی",   # عربی → فارسی
        "ة": "ه",
        "ؤ": "و",
        "إ": "ا",
        "أ": "ا",
        "آ": "ا",   # گاهی کاربران آ را بدون همزه می‌نویسند
    }

    def _normalize(self, text: str) -> str:
        """اصلاح حروف عربی/فارسی."""
        for wrong, correct in self.PERSIAN_CORRECTIONS.items():
            text = text.replace(wrong, correct)
        return text.strip()

    def get(self, request):
        query = request.query_params.get("q", "").strip()
        if len(query) < 2:
            return Response({"results": [], "query": query, "corrected": False})

        normalized = self._normalize(query)
        corrected = normalized != query

        from django.db.models import Q
        products = (
            Product.objects.filter(is_active=True)
            .filter(
                Q(name__icontains=normalized)
                | Q(brand__icontains=normalized)
                | Q(model_number__icontains=normalized)
                | Q(chipset__icontains=normalized)
                | Q(short_description__icontains=normalized)
            )
            .select_related("category")
            .prefetch_related(
                Prefetch("images", queryset=ProductImage.objects.filter(is_primary=True)),
                Prefetch("variants", queryset=ProductVariant.objects.filter(is_active=True)),
            )[:20]
        )

        serializer = ProductListSerializer(products, many=True, context={"request": request})
        return Response({
            "results": serializer.data,
            "count": len(serializer.data),
            "query": query,
            "normalized_query": normalized,
            "corrected": corrected,
        })


# ─── Exchange Rate ─────────────────────────────────────────────────────────────

class ExchangeRateView(APIView):
    """
    GET  /api/catalog/exchange-rate/   → نرخ ارز جاری (عموم)
    PUT  /api/catalog/exchange-rate/   → بروزرسانی (فقط ادمین)
    """

    def get_permissions(self):
        if self.request.method == "GET":
            return [AllowAny()]
        return [IsAdminUser()]

    def get(self, request):
        instance = ExchangeRateSetting.get_instance()
        serializer = ExchangeRateSerializer(instance)
        return Response(serializer.data)

    def put(self, request):
        instance = ExchangeRateSetting.get_instance()
        serializer = ExchangeRateUpdateSerializer(
            instance, data=request.data, context={"request": request}
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()
        logger.info(f"Exchange rate updated to {instance.rate_cny_to_irt} by {request.user}")
        return Response(ExchangeRateSerializer(instance).data)


# ─── Bulk Upload (Excel) ──────────────────────────────────────────────────────

@api_view(["POST"])
@permission_classes([IsAdminUser])
def bulk_upload_products(request):
    """
    POST /api/catalog/products/bulk-upload/
    Content-Type: multipart/form-data
    فیلد: file (xlsx)

    فرمت اکسل موردانتظار:
    sku | product_name | category_slug | chipset | voltage |
    lens_mount | part_type | purchase_price_usd | shipping_cost |
    profit_margin | stock_quantity
    """
    import openpyxl
    from decimal import Decimal, InvalidOperation

    file = request.FILES.get("file")
    if not file:
        return Response({"error": "فایل اکسل ارسال نشد."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
    except Exception as e:
        return Response({"error": f"خطا در خواندن فایل: {e}"}, status=status.HTTP_400_BAD_REQUEST)

    results = {"created": 0, "updated": 0, "errors": []}
    headers = None

    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_idx == 1:
            headers = [str(h).strip().lower() if h else "" for h in row]
            continue

        row_data = dict(zip(headers, row))
        sku = str(row_data.get("sku", "")).strip()
        if not sku:
            continue

        try:
            # دسته‌بندی
            cat_slug = str(row_data.get("category_slug", "")).strip()
            try:
                category = Category.objects.get(slug=cat_slug)
            except Category.DoesNotExist:
                results["errors"].append(f"ردیف {row_idx}: دسته '{cat_slug}' یافت نشد")
                continue

            # Product
            product_name = str(row_data.get("product_name", "")).strip()
            product, _ = Product.objects.get_or_create(
                name=product_name,
                defaults={
                    "category": category,
                    "part_type": row_data.get("part_type", "other"),
                    "chipset": str(row_data.get("chipset", "") or ""),
                    "voltage": str(row_data.get("voltage", "") or ""),
                    "lens_mount": row_data.get("lens_mount", "none"),
                },
            )

            # ProductVariant
            purchase_usd = Decimal(str(row_data.get("purchase_price_usd", 0) or 0))
            variant, created = ProductVariant.objects.update_or_create(
                sku=sku,
                defaults={
                    "product": product,
                    "purchase_price_usd": purchase_usd,
                    "shipping_cost": int(row_data.get("shipping_cost", 0) or 0),
                    "profit_margin": int(row_data.get("profit_margin", 0) or 0),
                    "stock_quantity": int(row_data.get("stock_quantity", 0) or 0),
                },
            )
            if created:
                results["created"] += 1
            else:
                results["updated"] += 1

        except (ValueError, InvalidOperation) as e:
            results["errors"].append(f"ردیف {row_idx} (SKU: {sku}): {e}")

    return Response(results, status=status.HTTP_200_OK)


# ─── Variants (for Kardex) ───────────────────────────────────

class VariantListView(generics.ListAPIView):
    """
    GET /api/catalog/variants/
    لیست همه variant ها برای استفاده در صفحه کاردکس
    """
    permission_classes = [IsAdminUser]

    def get(self, request):
        from .models import ProductVariant
        qs = ProductVariant.objects.filter(is_active=True).select_related('product', 'product__category')
        from apps.inventory.models import WarehouseStock
        from django.db.models import Sum

        # موجودی واقعی از WarehouseStock
        stock_map = {}
        for ws in WarehouseStock.objects.filter(warehouse__is_active=True).values('variant_id').annotate(total=Sum('quantity')):
            stock_map[ws['variant_id']] = ws['total']

        data = []
        for v in qs:
            # موجودی واقعی: اول از WarehouseStock، اگر نبود از stock_quantity
            real_stock = stock_map.get(v.id, v.stock_quantity)
            data.append({
                'id': v.id,
                'sku': v.sku,
                'name': v.product.name + (' - ' + v.name_suffix if v.name_suffix else ''),
                'product_name': v.product.name,
                'product_id': v.product.id,
                'category': v.product.category.name if v.product.category else '',
                'stock_quantity': real_stock,
            })
        return Response(data)


class SimilarProductsView(APIView):
    """GET /api/catalog/products/<slug>/similar/ — محصولات مشابه برای کالای ناموجود"""
    permission_classes = [AllowAny]

    def get(self, request, slug):
        from django.db.models import Q
        try:
            product = Product.objects.get(slug=slug, is_active=True)
        except Product.DoesNotExist:
            return Response([])

        # بررسی موجودی
        total_stock = sum(v.stock_quantity for v in product.variants.filter(is_active=True))

        # همیشه محصولات مشابه برمی‌گردانیم (حتی اگر موجود باشد)
        filters = Q(is_active=True) & ~Q(id=product.id)

        # فیلتر بر اساس مشخصات مشترک
        conditions = []
        if product.sensor_name:
            conditions.append(Q(sensor_name=product.sensor_name))
        if product.voltage:
            conditions.append(Q(voltage=product.voltage))
        if product.category:
            conditions.append(Q(category=product.category))

        similar = Product.objects.none()

        # اول: همان سنسور و ولتاژ
        if product.sensor_name and product.voltage:
            similar = Product.objects.filter(
                filters & Q(sensor_name=product.sensor_name) & Q(voltage=product.voltage)
            ).prefetch_related('variants', 'images')[:3]

        # اگر کافی نبود: همان دسته‌بندی
        if similar.count() < 3 and product.category:
            exclude_ids = [p.id for p in similar] + [product.id]
            extra = Product.objects.filter(
                filters & Q(category=product.category)
            ).exclude(id__in=exclude_ids).prefetch_related('variants', 'images')[:3 - similar.count()]
            from itertools import chain
            similar = list(chain(similar, extra))

        result = []
        for p in similar[:3]:
            variants = p.variants.filter(is_active=True)
            stock = sum(v.stock_quantity for v in variants)
            first_image = p.images.first()
            result.append({
                'id': p.id,
                'name': p.name,
                'slug': p.slug,
                'category': p.category.name if p.category else '',
                'sensor_name': p.sensor_name or '',
                'voltage': p.voltage or '',
                'stock': stock,
                'is_available': stock > 0,
                'image': first_image.image.url if first_image else None,
                'price': variants.first().final_price if variants.exists() else 0,
            })

        return Response({
            'is_out_of_stock': total_stock == 0,
            'similar': result,
        })


class PackageListView(APIView):
    """GET /api/catalog/packages/ — لیست پکیج‌ها"""
    permission_classes = [AllowAny]

    def get(self, request):
        from .models import Package
        packages = Package.objects.filter(is_active=True).prefetch_related(
            'package_items__product__images',
            'package_items__variant',
        )
        featured_only = request.query_params.get('featured') == '1'
        if featured_only:
            packages = packages.filter(is_featured=True)

        result = []
        for pkg in packages:
            items = []
            for item in pkg.package_items.all():
                img = item.product.images.first()
                items.append({
                    'product_id': item.product_id,
                    'product_name': item.product.name,
                    'product_slug': item.product.slug,
                    'variant_id': item.variant_id,
                    'quantity': item.quantity,
                    'unit_price': item.variant.final_price,
                    'image': img.image.url if img else None,
                })
            result.append({
                'id': pkg.id,
                'name': pkg.name,
                'slug': pkg.slug,
                'description': pkg.description,
                'image': pkg.image.url if pkg.image else None,
                'discount_percent': float(pkg.discount_percent),
                'original_price': pkg.original_price,
                'final_price': pkg.final_price,
                'savings': pkg.savings,
                'is_featured': pkg.is_featured,
                'items': items,
            })
        return Response(result)


class PackageDetailView(APIView):
    """GET /api/catalog/packages/<slug>/ — جزئیات پکیج"""
    permission_classes = [AllowAny]

    def get(self, request, slug):
        from .models import Package
        try:
            pkg = Package.objects.prefetch_related(
                'package_items__product__images',
                'package_items__variant',
            ).get(slug=slug, is_active=True)
        except Package.DoesNotExist:
            return Response({'error': 'یافت نشد'}, status=404)

        items = []
        for item in pkg.package_items.all():
            img = item.product.images.first()
            items.append({
                'product_id': item.product_id,
                'product_name': item.product.name,
                'product_slug': item.product.slug,
                'variant_id': item.variant_id,
                'quantity': item.quantity,
                'unit_price': item.variant.final_price,
                'subtotal': item.variant.final_price * item.quantity,
                'image': img.image.url if img else None,
                'stock': item.variant.stock_quantity,
            })

        return Response({
            'id': pkg.id,
            'name': pkg.name,
            'slug': pkg.slug,
            'description': pkg.description,
            'image': pkg.image.url if pkg.image else None,
            'discount_percent': float(pkg.discount_percent),
            'original_price': pkg.original_price,
            'final_price': pkg.final_price,
            'savings': pkg.savings,
            'items': items,
        })


class ProductPriceHistoryView(APIView):
    """GET /api/catalog/products/<slug>/price-history/ — تاریخچه قیمت ۳۰ روزه"""
    permission_classes = [AllowAny]

    def get(self, request, slug):
        from datetime import date, timedelta
        from .models import ProductPriceHistory, Product
        try:
            product = Product.objects.get(slug=slug)
        except Product.DoesNotExist:
            return Response([], status=404)

        variant = product.variants.filter(is_active=True).first()
        if not variant:
            return Response([])

        # ذخیره قیمت امروز
        today = date.today()
        ProductPriceHistory.objects.get_or_create(
            variant=variant, date=today,
            defaults={'price_irr': variant.final_price, 'price_cny': variant.purchase_price_cny or 0}
        )

        # ۳۰ روز اخیر
        since = today - timedelta(days=30)
        history = ProductPriceHistory.objects.filter(
            variant=variant, date__gte=since
        ).order_by('date')

        return Response([{
            'date': str(h.date),
            'price': h.price_irr,
        } for h in history])
