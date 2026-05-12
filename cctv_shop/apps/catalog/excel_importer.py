"""
apps/catalog/excel_importer.py
وارد کردن محصولات از Excel
"""
import openpyxl
from django.db import transaction
from apps.catalog.models import Category, Product, ProductVariant


# نگاشت نام ستون‌ها
COLUMNS = {
    # اطلاعات اصلی
    "نام محصول":       "name",
    "دسته بندی":       "category_slug",
    "شرکت سازنده":     "brand",
    "شماره مدل":       "model_number",
    "توضیح کوتاه":     "short_description",
    "توضیحات":         "description",
    # قیمت
    "قیمت خرید (یوان)":  "purchase_price_cny",
    "هزینه حمل (یوان)":  "shipping_cost_cny",
    "سود (%)":            "profit_percent",
    "موجودی":             "stock_quantity",
    "کد SKU":             "sku",
    # برد
    "نام برد":          "board_name",
    "سیستم":            "board_system",
    "رزولوشن برد":      "board_resolution",
    "نام سنسور":        "sensor_name",
    "اندازه سنسور":     "sensor_size",
    "AI پردازنده":      "ai_processor",
    "ابعاد برد":        "board_dimensions",
    # لنز
    "نام لنز":          "lens_name",
    "فاصله کانونی":     "lens_degree",
    "نوع لنز":          "lens_type",
    "نوع کاربرد":       "focal_ratio",
    "رزولوشن لنز":      "lens_resolution",
    # متعلقات
    "مدل":              "modem_model",
    "نوع":              "modem_type",
    "تعداد پین":        "pin_count",
    "ولتاژ":            "voltage",
    "آمپر خروجی":       "ampere_output",
    "قابلیت سیم کارت": "sim_card",
    "تعداد پورت":       "output_ports",
    "سایز":             "modem_size",
    # بدنه
    "جنس بدنه":         "housing_material",
    "نوع بدنه":         "housing_type",
    "سایز دوربین":      "housing_size",
    "قابلیت میکروفون": "has_microphone",
}

BOOL_FIELDS = {"sim_card", "has_microphone"}
INT_FIELDS  = {"stock_quantity", "output_ports"}
DEC_FIELDS  = {"purchase_price_cny", "shipping_cost_cny", "profit_percent"}


def import_products_from_excel(file_path: str, dry_run: bool = False) -> dict:
    wb = openpyxl.load_workbook(file_path, data_only=True)

    # پیدا کردن sheet محصولات
    ws = None
    for name in ["محصولات", "Products", "products", wb.sheetnames[0]]:
        if name in wb.sheetnames:
            ws = wb[name]
            break
    if not ws:
        ws = wb.active

    results = {"created": 0, "updated": 0, "errors": []}

    # خواندن هدر
    headers = {}
    for col in ws.iter_cols(1, ws.max_column, 1, 1):
        cell = col[0]
        if cell.value and str(cell.value).strip() in COLUMNS:
            headers[cell.column] = COLUMNS[str(cell.value).strip()]

    if not headers:
        results["errors"].append("ستون‌های مورد نظر در فایل پیدا نشد")
        return results

    # cache دسته‌بندی‌ها
    cat_cache = {c.slug: c for c in Category.objects.filter(is_active=True)}
    cat_name_cache = {c.name: c for c in Category.objects.filter(is_active=True)}

    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(row):
            continue

        # ساخت dict از ردیف
        data = {}
        for col_idx, field in headers.items():
            val = row[col_idx - 1]
            if val is not None:
                data[field] = val

        if not data.get("name"):
            continue

        try:
            with transaction.atomic():
                # پیدا کردن دسته‌بندی
                cat = None
                cat_val = str(data.pop("category_slug", "") or "").strip()
                if cat_val:
                    cat = cat_cache.get(cat_val) or cat_name_cache.get(cat_val)

                if not cat:
                    results["errors"].append(f"ردیف {row_idx}: دسته‌بندی '{cat_val}' پیدا نشد")
                    continue

                # اطلاعات variant
                sku               = str(data.pop("sku", "") or "").strip()
                purchase_price_cny= data.pop("purchase_price_cny", 0) or 0
                shipping_cost_cny = data.pop("shipping_cost_cny", 0) or 0
                profit_percent    = data.pop("profit_percent", 0) or 0
                stock_quantity    = int(data.pop("stock_quantity", 0) or 0)

                # تبدیل نوع‌ها
                for field in list(data.keys()):
                    val = data[field]
                    if field in BOOL_FIELDS:
                        data[field] = str(val).strip() in ("1","بله","yes","true","دارد","✓")
                    elif field in INT_FIELDS:
                        try: data[field] = int(val)
                        except: data[field] = None
                    elif field in DEC_FIELDS:
                        try: data[field] = float(val)
                        except: data[field] = 0
                    elif val is not None:
                        data[field] = str(val).strip()

                # ساخت یا بروزرسانی محصول
                if not dry_run:
                    product, created = Product.objects.update_or_create(
                        name=data["name"],
                        defaults={**data, "category": cat},
                    )
                    if created:
                        results["created"] += 1
                    else:
                        results["updated"] += 1

                    # ساخت یا بروزرسانی variant
                    if not sku:
                        sku = f"AUTO-{product.pk}"
                    ProductVariant.objects.update_or_create(
                        sku=sku,
                        defaults={
                            "product": product,
                            "purchase_price_cny": purchase_price_cny,
                            "shipping_cost_cny":  shipping_cost_cny,
                            "profit_percent":     profit_percent,
                            "stock_quantity":     stock_quantity,
                            "is_active": True,
                        }
                    )
                else:
                    results["created"] += 1

        except Exception as e:
            results["errors"].append(f"ردیف {row_idx}: {e}")

    return results
